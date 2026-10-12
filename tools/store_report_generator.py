#!/usr/bin/env python3
"""
Store Summary Report Generator - Flask Blueprint
Mounted at /store-report by main.py. Not meant to be run directly.
"""

import sys
import os
import re
import json
from io import BytesIO
from pathlib import Path

from flask import Blueprint, render_template_string, request, send_file, jsonify
from openpyxl import load_workbook
from docx import Document

# tools/store_report_generator.py -> project_root/data/template.docx
project_root = Path(__file__).parent.parent
template_path = project_root / "data" / "template.docx"

# HTML Form (UI for the web interface)
HTML_FORM = """<!DOCTYPE html><html><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1.0"><title>Store Summary Report Generator</title><style>* { margin: 0; padding: 0; box-sizing: border-box; }body { font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; background-color: #c62b27; background-image: url('https://fundraising.hotdogonastick.com/_next/static/media/section-bg.3d099ea1.jpg'); background-size: cover; background-attachment: scroll; min-height: 100vh; padding: 0; }.container { max-width: 1000px; margin: 0 auto 0; background: white; border-radius: 12px 12px 0 0; box-shadow: 0 20px 60px rgba(0,0,0,0.3); }.header { background: #e53430; color: white; padding: 40px; text-align: center; }.header h1 { font-size: 2.5em; margin-bottom: 10px; }.header p { font-size: 1.1em; opacity: 0.9; }.content { padding: 40px; }.form-group { margin-bottom: 20px; }label { display: block; font-weight: 600; margin-bottom: 8px; color: #333; }input[type="text"], input[type="number"], input[type="file"], input[type="date"], textarea, select { width: 100%; padding: 12px; border: 2px solid #ddd; border-radius: 6px; font-size: 1em; font-family: inherit; }input:focus, textarea:focus, select:focus { outline: none; border-color: #0063a9; box-shadow: 0 0 0 3px rgba(0, 99, 169, 0.1); }textarea { resize: vertical; min-height: 80px; }.two-column { display: grid; grid-template-columns: 1fr 1fr; gap: 20px; }.button-group { display: flex; gap: 10px; margin-top: 30px; }button { flex: 1; padding: 14px 28px; font-size: 1em; font-weight: 600; border: none; border-radius: 6px; cursor: pointer; transition: all 0.3s; }.btn-primary { background: #e53430; color: white; }.btn-primary:hover { transform: translateY(-2px); box-shadow: 0 10px 20px rgba(229, 52, 48, 0.3); }.btn-primary:disabled { opacity: 0.6; cursor: not-allowed; }.btn-secondary { background: #f0f0f0; color: #333; }.btn-secondary:hover { background: #e0e0e0; }.status { padding: 15px; border-radius: 6px; margin-bottom: 20px; display: none; }.status.success { background: #d4edda; color: #155724; border: 1px solid #c3e6cb; }.status.error { background: #f8d7da; color: #721c24; border: 1px solid #f5c6cb; }.status.loading { background: #d1ecf1; color: #0c5460; border: 1px solid #bee5eb; }.repair-row { display: grid; grid-template-columns: 1fr 2fr 1fr 2fr; gap: 10px; margin-bottom: 10px; padding: 15px; border: 1px solid #ddd; border-radius: 6px; }.repair-row button { padding: 8px; font-size: 0.9em; }@media (max-width: 600px) { .two-column { grid-template-columns: 1fr; } .button-group { flex-direction: column; } .repair-row { grid-template-columns: 1fr; } }</style></head><body><div style="width:100%;height:40px;background:repeating-linear-gradient(90deg,#e53430 0px,#e53430 60px,#0063a9 60px,#0063a9 120px,#fed500 120px,#fed500 180px,#f1f1f1 180px,#f1f1f1 240px);"></div><div class="container"><div class="header"><h1>Store Summary Report Generator</h1></div><div class="content"><div id="status" class="status"></div><form id="reportForm"><div class="two-column"><div class="form-group"><label for="currentFile">Current Week Excel *</label><input type="file" id="currentFile" accept=".xlsx" required></div><div class="form-group"><label for="lastYearFile">Last Year Excel *</label><input type="file" id="lastYearFile" accept=".xlsx" required></div></div><div class="two-column"><div class="form-group"><label for="budgetFile">Budget Excel</label><input type="file" id="budgetFile" accept=".xlsx"></div><div class="form-group"><label for="budgetPeriod">Budget Period</label><select id="budgetPeriod"><option value="Period 1">Period 1</option><option value="Period 2">Period 2</option><option value="Period 3">Period 3</option><option value="Period 4">Period 4</option><option value="Period 5">Period 5</option><option value="Period 6">Period 6</option><option value="Period 7">Period 7</option><option value="Period 8">Period 8</option><option value="Period 9">Period 9</option><option value="Period 10">Period 10</option><option value="Period 11">Period 11</option><option value="Period 12">Period 12</option></select></div></div><div class="two-column"><div class="form-group"><label for="storeNameNumber">Store Name/Number *</label><input type="text" id="storeNameNumber" placeholder="e.g., Thousand Oaks #36" required></div><div class="form-group"><label for="managerName">Manager Name *</label><input type="text" id="managerName" placeholder="e.g., Daniel Nehemne" required></div></div><div class="two-column"><div class="form-group"><label for="hoursScheduled">Hours Scheduled</label><input type="number" id="hoursScheduled" placeholder="e.g., 78" step="0.01"></div><div class="form-group"><label for="salesBudget">Sales Budget</label><input type="number" id="salesBudget" placeholder="e.g., 5000.00" step="0.01"></div></div><div class="two-column"><div class="form-group"><label for="prevSummaryFile">Previous Summary (.docx)</label><input type="file" id="prevSummaryFile" accept=".docx"></div><div class="form-group"><label for="currentCogs">This Week COGS $</label><input type="number" id="currentCogs" placeholder="e.g., 2000.00" step="0.01"></div></div><div class="form-group"><label for="weekendingDate">Weekending Date *</label><input type="text" id="weekendingDate" placeholder="Auto-filled from filename" readonly></div><div class="form-group"><label for="salesComments">Sales Comments</label><textarea id="salesComments" placeholder="Enter sales comments..."></textarea></div><div class="form-group"><label for="laborComments">Labor Comments</label><textarea id="laborComments" placeholder="Enter labor comments..."></textarea></div><div class="form-group"><label for="financialComments">Financial Comments</label><textarea id="financialComments" placeholder="Enter financial comments..."></textarea></div><div class="form-group"><label for="teamComments">Team Update Comments</label><textarea id="teamComments" placeholder="Enter team update comments..."></textarea></div><div class="form-group"><label>86 Repairs</label><div id="repairsContainer"></div><button type="button" class="btn-secondary" style="margin-top: 10px; width: 100%;" onclick="addRepairRow()">+ Add Repair</button></div><div class="button-group"><button type="submit" class="btn-primary" id="submitBtn">Generate Report</button><button type="button" class="btn-secondary" onclick="location.reload();">Clear Form</button></div></form></div></div><script>function showStatus(message, type) { const status = document.getElementById('status'); status.textContent = message; status.className = `status ${type}`; status.style.display = 'block'; }function addRepairRow() { const container = document.getElementById('repairsContainer'); const row = document.createElement('div'); row.className = 'repair-row'; row.innerHTML = `<input type="date" class="repair-date" value=""><input type="text" placeholder="Repair Description" class="repair-desc" value=""><input type="text" placeholder="Status" class="repair-status" value=""><input type="text" placeholder="Notes" class="repair-notes" value=""><button type="button" class="btn-secondary" style="grid-column: 1/-1;" onclick="this.parentElement.remove()">Remove</button>`; container.appendChild(row); }addRepairRow();function extractDateFromFilename(filename) { const matches = [...filename.matchAll(/(\\d{4})-(\\d{2})-(\\d{2})/g)]; if (matches.length > 0) { const lastMatch = matches[matches.length - 1]; return `${lastMatch[2]}/${lastMatch[3]}/${lastMatch[1].slice(2)}`; } return new Date().toLocaleDateString('en-US'); }document.getElementById('currentFile').addEventListener('change', (e) => { const file = e.target.files[0]; if (file) { const date = extractDateFromFilename(file.name); document.getElementById('weekendingDate').value = date; } });document.getElementById('reportForm').addEventListener('submit', async (e) => { e.preventDefault(); try { showStatus('Processing files...', 'loading'); document.getElementById('submitBtn').disabled = true; const currentFile = document.getElementById('currentFile').files[0]; const lastYearFile = document.getElementById('lastYearFile').files[0]; if (!currentFile || !lastYearFile) { showStatus('Please select both Excel files', 'error'); document.getElementById('submitBtn').disabled = false; return; } const formData = new FormData(); formData.append('currentFile', currentFile); formData.append('lastYearFile', lastYearFile); const budgetFile = document.getElementById('budgetFile').files[0]; if (budgetFile) { formData.append('budgetFile', budgetFile); } const prevSummaryFile = document.getElementById('prevSummaryFile').files[0]; if (prevSummaryFile) { formData.append('prevSummaryFile', prevSummaryFile); } formData.append('budgetPeriod', document.getElementById('budgetPeriod').value); formData.append('storeNameNumber', document.getElementById('storeNameNumber').value); formData.append('managerName', document.getElementById('managerName').value); formData.append('weekendingDate', document.getElementById('weekendingDate').value); formData.append('hoursScheduled', document.getElementById('hoursScheduled').value || 0); formData.append('salesBudget', document.getElementById('salesBudget').value || 0); formData.append('currentCogs', document.getElementById('currentCogs').value || 0); formData.append('salesComments', document.getElementById('salesComments').value); formData.append('laborComments', document.getElementById('laborComments').value); formData.append('financialComments', document.getElementById('financialComments').value); formData.append('teamComments', document.getElementById('teamComments').value); const repairs = []; document.querySelectorAll('#repairsContainer > div').forEach(row => { let date = row.querySelector('.repair-date').value; const desc = row.querySelector('.repair-desc').value; const status = row.querySelector('.repair-status').value; const notes = row.querySelector('.repair-notes').value; if (date) { const [year, month, day] = date.split('-'); date = `${month}/${day}/${year.slice(2)}`; } if (date || desc || status || notes) { repairs.push({ date, desc, status, notes }); } }); formData.append('repairs', JSON.stringify(repairs)); showStatus('Generating report...', 'loading'); const response = await fetch('/store-report/generate', { method: 'POST', body: formData }); if (!response.ok) { const error = await response.json(); showStatus(`Error: ${error.error}`, 'error'); document.getElementById('submitBtn').disabled = false; return; } const blob = await response.blob(); const url = URL.createObjectURL(blob); const link = document.createElement('a'); link.href = url; link.download = `Store_Summary_Report_${document.getElementById('weekendingDate').value.replace(/\\//g, '-')}.docx`; link.click(); URL.revokeObjectURL(url); showStatus('Report generated!', 'success'); document.getElementById('submitBtn').disabled = false; } catch (error) { console.error('Error:', error); showStatus(`Error: ${error.message}`, 'error'); document.getElementById('submitBtn').disabled = false; } });</script><div style="width:100%;height:40px;background:repeating-linear-gradient(90deg,#e53430 0px,#e53430 60px,#0063a9 60px,#0063a9 120px,#fed500 120px,#fed500 180px,#f1f1f1 180px,#f1f1f1 240px);"></div></body></html>"""

def iter_paragraphs_in_order(document):
    """Yield paragraphs in true document order, including those inside table cells."""
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    for child in document.element.body:
        if child.tag.endswith('}p'):
            yield Paragraph(child, document)
        elif child.tag.endswith('}tbl'):
            tbl = Table(child, document)
            for row in tbl.rows:
                for cell in row.cells:
                    yield from cell.paragraphs

bp = Blueprint('store_report', __name__, url_prefix='/store-report')

@bp.route('/')
def index():
    return render_template_string(HTML_FORM)

@bp.route('/generate', methods=['POST'])
def generate():
    print("=== GENERATE ENDPOINT HIT ===")
    try:
        # Get files
        current_file = request.files['currentFile']
        last_year_file = request.files['lastYearFile']
        
        # Extract data from Excel files
        current_metrics = extract_excel_data(current_file)
        last_year_metrics = extract_excel_data(last_year_file)

        # Load template
        doc = Document(str(template_path))

        # Update store info table
        if len(doc.tables) > 0:
            doc.tables[0].rows[0].cells[0].text = f"Store Name/#: {request.form['storeNameNumber']}\nManager Completing Report: {request.form['managerName']}"
            doc.tables[0].rows[1].cells[0].text = f"Weekending Date: {request.form['weekendingDate']}"

        # Update sales table
        if len(doc.tables) > 1:
            write_cell_value(doc.tables[1].rows[1].cells[0], f"${current_metrics['sales']:.2f}")
            write_cell_value(doc.tables[1].rows[1].cells[1], str(current_metrics['transactions']))
            write_cell_value(doc.tables[1].rows[1].cells[2], f"${current_metrics['avgTicket']}")
            write_cell_value(doc.tables[1].rows[3].cells[0], f"${last_year_metrics['sales']:.2f}")
            write_cell_value(doc.tables[1].rows[3].cells[1], str(last_year_metrics['transactions']))
            write_cell_value(doc.tables[1].rows[3].cells[2], f"${last_year_metrics['avgTicket']}")

        # Update labor table
        hours_worked = current_metrics['hoursWorked']
        if len(doc.tables) > 2:
            write_cell_value(doc.tables[2].rows[1].cells[0], str(hours_worked))
            write_cell_value(doc.tables[2].rows[1].cells[1], str(float(request.form.get('hoursScheduled', 0))))
            diff = hours_worked - float(request.form.get('hoursScheduled', 0))
            write_cell_value(doc.tables[2].rows[1].cells[2], f"{diff:.2f}")

        # Update budget and actual rows (Table 3)
        t3 = doc.tables[3] if len(doc.tables) > 3 else None
        _budget_vals = None  # saved for difference row
        _actual_vals = None  # saved for difference row

        # Budget row (rows[0]) — from budget Excel
        budget_file = request.files.get('budgetFile')
        if t3 and budget_file:
            budget = extract_budget_data(budget_file, request.form.get('budgetPeriod'),
                                        request.form.get('storeNameNumber'))
            if budget:
                write_budget_cell(t3.rows[0].cells[0], f"{budget['salesBudget']:.2f}")
                write_budget_cell(t3.rows[0].cells[1], f"{budget['cogsBudgetPct']}% ${budget['cogsBudgetDollars']:.2f}")
                write_budget_cell(t3.rows[0].cells[2], f"{budget['laborBudgetPct']}%")
                _budget_vals = budget

        # Actual row (rows[1]) — cumulative period-to-date
        _debug_path = str(project_root / "debug.txt")

        if t3 and len(t3.rows) > 1:
            current_sales = float(current_metrics['sales'])
            current_cogs  = float(request.form.get('currentCogs', 0) or 0)
            # Labor cost is auto-extracted from Excel
            current_labor = float(current_metrics.get('laborCost', 0))

            prior = {'priorSales': 0.0, 'priorCogs': 0.0, 'priorLabor': 0.0}
            prev_file = request.files.get('prevSummaryFile')
            if prev_file and prev_file.filename:
                parsed = extract_actual_from_summary(prev_file)
                if parsed:
                    prior = parsed

            total_sales = prior['priorSales'] + current_sales
            total_cogs  = prior['priorCogs']  + current_cogs
            total_labor = prior['priorLabor'] + current_labor

            if total_sales > 0:
                raw_labor_pct = total_labor / total_sales * 100
                cogs_pct  = format_percent(total_cogs  / total_sales * 100)
                labor_pct = format_percent(raw_labor_pct)
            else:
                cogs_pct = labor_pct = "0"

            with open(_debug_path, "w") as f:
                f.write(f"current_sales={current_sales}\n")
                f.write(f"current_labor={current_labor}\n")
                f.write(f"current_cogs={current_cogs}\n")
                f.write(f"total_sales={total_sales}\n")
                f.write(f"total_labor={total_labor}\n")
                f.write(f"total_cogs={total_cogs}\n")
                f.write(f"raw_labor_pct={total_labor / total_sales * 100 if total_sales else 'N/A'}\n")
                f.write(f"labor_pct_string='{labor_pct}'\n")
                f.write(f"cogs_pct_string='{cogs_pct}'\n")
            sys.stderr.write(f"[DEBUG] labor_pct='{labor_pct}' total_labor={total_labor} total_sales={total_sales}\n")
            sys.stderr.flush()

            write_budget_cell(t3.rows[1].cells[0], f"${total_sales:.2f}")
            write_budget_cell(t3.rows[1].cells[1], f"{cogs_pct}% ${total_cogs:.2f}")
            write_budget_cell(t3.rows[1].cells[2], f"{labor_pct}% ${total_labor:.2f}")
            _actual_vals = {
                'sales': total_sales,
                'cogs': total_cogs,
                'labor_pct': float(labor_pct),
            }
        else:
            with open(_debug_path, "w") as f:
                f.write(f"t3={t3}\n")
                f.write(f"t3 is None: {t3 is None}\n")
                if t3:
                    f.write(f"t3.rows length: {len(t3.rows)}\n")

        # Difference row (rows[3]) — actual minus budget
        # Sales/COGS: dollar difference. Labor: percentage point difference.
        if t3 and len(t3.rows) > 3 and _budget_vals and _actual_vals:
            sales_diff = _actual_vals['sales'] - _budget_vals['salesBudget']
            cogs_diff  = _actual_vals['cogs']  - _budget_vals['cogsBudgetDollars']
            labor_diff = _actual_vals['labor_pct'] - float(_budget_vals['laborBudgetPct'])

            def fmt_dollar_diff(v):
                return f"-${abs(v):.2f}" if v < 0 else f"${v:.2f}"

            def fmt_pct_diff(v):
                s = format_percent(abs(v))
                return f"-{s}%" if v < 0 else f"{s}%"

            write_budget_cell(t3.rows[3].cells[0], fmt_dollar_diff(sales_diff))
            write_budget_cell(t3.rows[3].cells[1], fmt_dollar_diff(cogs_diff))
            write_budget_cell(t3.rows[3].cells[2], fmt_pct_diff(labor_diff))

        # Write financial fields — match by label text in whichever table they appear in
        fill_word_label_cell(doc.tables, 'Cash Variance',  f"${current_metrics['cashVariance']:.2f}")
        fill_word_label_cell(doc.tables, 'Cleared Items',  str(int(current_metrics['clearedItems'])))
        fill_word_label_cell(doc.tables, 'No Sales',       str(int(current_metrics['noSales'])))
        fill_word_label_cell(doc.tables, 'Online Sales',   f"${current_metrics['onlineSales']:.2f}")
        fill_word_label_cell(doc.tables, 'Refunds',        f"${current_metrics['refunds']:.2f}")
        fill_word_label_cell(doc.tables, 'Payout/Payin',
                             f"In: ${current_metrics['payIns']:.2f} / Out: ${current_metrics['payOuts']:.2f}")
        
        # Update section comments.
        # Use doc.paragraphs (body-level only) — NOT iter_paragraphs_in_order.
        # Reason: table cells inside the Financials table contain "Labor Budget %",
        # which would falsely reset current_section back to 'Labor' before the
        # Financial Comments paragraph is reached if we iterate table cells too.
        current_section = None
        for para in doc.paragraphs:
            t = para.text.strip()
            if not t:
                continue
            if 'Comments' not in t and 'Repair' not in t:
                if 'Sales' in t and 'Last Year' not in t and 'Budget' not in t:
                    current_section = 'Sales'
                elif 'Labor' in t and 'Budget' not in t:
                    current_section = 'Labor'
                elif 'Financial' in t:
                    current_section = 'Financial'
            if t.startswith('Comments:'):
                if current_section == 'Sales':
                    para.text = f"Comments: {request.form.get('salesComments', '')}"
                elif current_section == 'Labor':
                    para.text = f"Comments: {request.form.get('laborComments', '')}"
                elif current_section == 'Financial':
                    para.text = f"Comments: {request.form.get('financialComments', '')}"

        # Team Update: find the table (the box) that follows the "Team Update"
        # body paragraph and write the comment into it, clearing all placeholder rows.
        from docx.table import Table as DocxTable
        for para in doc.paragraphs:
            if 'Team Update' in para.text:
                found = False
                for child in doc.element.body:
                    if child.tag.endswith('}p') and child is para._element:
                        found = True
                    elif found and child.tag.endswith('}tbl'):
                        team_table = DocxTable(child, doc)
                        first_cell = team_table.rows[0].cells[0]
                        first_cell.paragraphs[0].text = request.form.get('teamComments', '')
                        for p in first_cell.paragraphs[1:]:
                            p.text = ''
                        for row in team_table.rows[1:]:
                            for cell in row.cells:
                                for p in cell.paragraphs:
                                    p.text = ''
                        break
                break

        # Update repairs table: template rows have "Date:" / "Repair:" / "Status:" /
        # "Notes:" as existing cell text — append the repair data after those labels.
        repairs = json.loads(request.form.get('repairs', '[]'))
        if len(doc.tables) > 5 and repairs:
            repair_table = doc.tables[5]
            repair_idx = 0
            for row in repair_table.rows:
                if repair_idx >= len(repairs):
                    break
                cells = row.cells
                if len(cells) >= 4 and cells[0].text.strip().startswith('Date'):
                    r = repairs[repair_idx]
                    write_cell_value(cells[0], f"Date: {r.get('date', '')}")
                    write_cell_value(cells[1], f"Repair: {r.get('desc', '')}")
                    write_cell_value(cells[2], f"Status: {r.get('status', '')}")
                    write_cell_value(cells[3], f"Notes: {r.get('notes', '')}")
                    repair_idx += 1
        
        # Save to bytes
        output = BytesIO()
        doc.save(output)
        output.seek(0)
        
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            as_attachment=True,
            download_name=f'Store_Summary_Report_{request.form["weekendingDate"].replace("/", "-")}.docx'
        )
        
    except Exception as e:
        print(f"ERROR IN GENERATE: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 400

@bp.route('/debug')
def debug():
    """Show the template's paragraph and table structure to help diagnose mapping issues."""
    doc = Document(str(template_path))
    lines = ['<pre style="font-family:monospace;font-size:13px;padding:20px">']
    lines.append('=== BODY + TABLE PARAGRAPHS (document order) ===\n')
    for i, para in enumerate(iter_paragraphs_in_order(doc)):
        t = repr(para.text[:120])
        lines.append(f'[para {i:03d}] {t}')
    lines.append('\n\n=== TABLES ===\n')
    for ti, tbl in enumerate(doc.tables):
        lines.append(f'\n-- Table {ti} ({len(tbl.rows)} rows x {len(tbl.columns)} cols) --')
        for ri, row in enumerate(tbl.rows):
            for ci, cell in enumerate(row.cells):
                txt = repr(cell.text[:80])
                lines.append(f'  [{ri},{ci}] {txt}')
    lines.append('</pre>')
    return '\n'.join(lines)

def extract_budget_data(file, period, store_number):
    """
    Extract budget values for a store from the budget Excel.
    Handles Excel formulas and decimal percent formats.
    Returns dict with salesBudget, laborBudgetDollars, laborBudgetPct,
    cogsBudgetDollars, cogsBudgetPct. Returns None if store not found.
    """
    try:
        if hasattr(file, 'seek'):
            file.seek(0)
        # Use data_only=True to read cached calculated values instead of formula strings
        wb = load_workbook(file, data_only=True)

        # Get the sheet for the selected period
        if period not in wb.sheetnames:
            return None
        ws = wb[period]

        # Extract store number from form input (e.g., "Thousand Oaks #36" -> 36)
        m = re.search(r'#(\d+)', store_number)
        if not m:
            return None
        store_num = int(m.group(1))

        # Find the row matching the store number
        mr = ws.max_row
        for row in range(1, mr + 1):
            cell_val = ws.cell(row=row, column=1).value
            try:
                cell_val_int = int(cell_val) if cell_val else None
                if cell_val_int and cell_val_int == store_num:
                    # Extract columns B-F (Sales $, Labor $, Labor %, COGS $, COGS %)
                    sales_budget = ws.cell(row=row, column=2).value or 0
                    labor_budget_dollars = ws.cell(row=row, column=3).value
                    labor_budget_pct = ws.cell(row=row, column=4).value or 0
                    cogs_budget_dollars = ws.cell(row=row, column=5).value
                    cogs_budget_pct = ws.cell(row=row, column=6).value or 0

                    # Convert to float and normalize decimal percents
                    sales_budget = float(sales_budget) if sales_budget else 0
                    labor_budget_pct = float(labor_budget_pct) if labor_budget_pct else 0
                    cogs_budget_pct = float(cogs_budget_pct) if cogs_budget_pct else 0

                    # Normalize percents: if < 1.5, multiply by 100 (convert 0.36 -> 36)
                    if labor_budget_pct and labor_budget_pct < 1.5:
                        labor_budget_pct *= 100
                    if cogs_budget_pct and cogs_budget_pct < 1.5:
                        cogs_budget_pct *= 100

                    # If formula cells returned None, compute from sales * percent
                    if not labor_budget_dollars and sales_budget:
                        labor_budget_dollars = sales_budget * (labor_budget_pct / 100)
                    if not cogs_budget_dollars and sales_budget:
                        cogs_budget_dollars = sales_budget * (cogs_budget_pct / 100)

                    budget_dict = {
                        'salesBudget': sales_budget,
                        'laborBudgetDollars': float(labor_budget_dollars) if labor_budget_dollars else 0,
                        'laborBudgetPct': format_percent(labor_budget_pct),
                        'cogsBudgetDollars': float(cogs_budget_dollars) if cogs_budget_dollars else 0,
                        'cogsBudgetPct': format_percent(cogs_budget_pct),
                    }
                    return budget_dict
            except (ValueError, TypeError):
                # Skip rows with non-numeric values (headers, etc.)
                continue
        return None
    except Exception as e:
        print(f"Error extracting budget data: {e}")
        import traceback
        traceback.print_exc()
        return None

def extract_actual_from_summary(file):
    """
    Parses prior period actuals from an uploaded previous week's .docx summary.
    Reads doc.tables[3].rows[1] (the actual row from that prior report).
    Extracts Sales $, COGS $, and Labor $ values.
    """
    try:
        from docx import Document as _Doc
        doc = _Doc(file)
        if len(doc.tables) <= 3 or len(doc.tables[3].rows) <= 1:
            return None
        cells = doc.tables[3].rows[1].cells
        def parse_dollar(text):
            # Extracts last dollar amount from e.g. "22% $6175.54" or "$28070.62"
            matches = re.findall(r'\$([\d,]+\.?\d*)', text)
            return float(matches[-1].replace(',', '')) if matches else 0.0
        return {
            'priorSales': parse_dollar(cells[0].text),
            'priorCogs':  parse_dollar(cells[1].text),
            'priorLabor': parse_dollar(cells[2].text),
        }
    except Exception as e:
        print(f"Error extracting prior summary: {e}")
        return None

def format_percent(value):
    """Format percentage to 2 decimals, removing trailing zeros."""
    return f"{float(value):.2f}".rstrip('0').rstrip('.')

def write_cell_value(cell, value_text):
    """
    Write value_text into a cell while preserving paragraph structure and formatting.
    Clears runs in the first paragraph and adds the new value, inheriting paragraph
    formatting (alignment, spacing) and default font settings.
    """
    if not cell or not cell.paragraphs:
        raise ValueError(f"Cell is invalid or has no paragraphs: {cell}")
    para = cell.paragraphs[0]
    # Remove all existing runs from the paragraph (preserves paragraph-level formatting)
    for run in para.runs[:]:
        r = run._r
        r.getparent().remove(r)
    # Add new run (inherits paragraph's formatting and defaults)
    para.add_run(value_text)

def write_budget_cell(cell, value_text):
    """
    Write value_text into a budget cell without disturbing the header paragraph.
    Leaves paragraph[0] (the label) untouched, writes value into paragraph[1].
    For cells with header + value layout.
    """
    paras = cell.paragraphs
    src_font = paras[0].runs[0].font if paras and paras[0].runs else None

    def apply_font(run):
        if src_font:
            run.font.name = src_font.name
            run.font.size = src_font.size
            run.font.bold = src_font.bold

    if len(paras) > 1:
        para = paras[1]
        for run in para.runs[:]:
            r = run._r
            r.getparent().remove(r)
        apply_font(para.add_run(value_text))
    else:
        apply_font(cell.add_paragraph().add_run(value_text))

def fill_word_label_cell(tables, label, value):
    """
    Find a table cell whose text contains label and write value into the next cell in
    the same row.  Matches both with and without spaces so 'Pay Ins' finds 'Payins'.
    Uses write_cell_value to preserve paragraph formatting.
    """
    label_lower   = label.lower()
    label_nospace = label_lower.replace(' ', '')
    for table in tables:
        for row in table.rows:
            for ci, cell in enumerate(row.cells):
                ct        = cell.text.lower()
                ct_nospace = ct.replace(' ', '')
                if (label_lower in ct or label_nospace in ct_nospace) \
                   and ci + 1 < len(row.cells):
                    write_cell_value(row.cells[ci + 1], str(value))
                    return True
    return False

LABOR_ROW_LABELS = {'Management', 'Service Crew'}

def extract_excel_data(file):
    """
    Label-based extraction of all metrics from an Excel operational report.
    Uses multiple targeted passes so section-specific logic doesn't interfere
    with simple label lookups, and so column headers (e.g. 'Qty') inside a
    section can't accidentally close that section.
    """
    _ZERO = {
        'sales': 0, 'transactions': 0, 'avgTicket': 0, 'hoursWorked': 0,
        'laborCost': 0, 'cashVariance': 0, 'clearedItems': 0, 'noSales': 0,
        'onlineSales': 0, 'refunds': 0, 'payIns': 0, 'payOuts': 0,
    }
    try:
        if hasattr(file, 'seek'):
            file.seek(0)
        wb = load_workbook(file)
        ws = wb.active
        mr = ws.max_row

        # ── Helpers ──────────────────────────────────────────────────────────
        def s(row, col=1):
            """Return cell value as a stripped string, '' if empty."""
            v = ws.cell(row=row, column=col).value
            return str(v).strip() if v is not None else ''

        def num(row, col):
            """Return cell value as float, None if not numeric."""
            v = ws.cell(row=row, column=col).value
            return float(v) if isinstance(v, (int, float)) else None

        def find_section(keyword):
            """Return the first row where col A contains keyword (case-insensitive)."""
            kw = keyword.upper()
            for r in range(1, mr + 1):
                if kw in s(r).upper():
                    return r
            return None

        def header_col(start_row, names):
            """
            Scan start_row and start_row+1 for a cell whose lowercased text is in
            names.  Returns that column number, or 2 (col B) as a fallback.
            """
            for r in (start_row, start_row + 1):
                for col in range(1, 20):
                    if s(r, col).lower() in names:
                        return col
            return 2

        def section_rows(start_row, skip=1, max_scan=60):
            """
            Yield row numbers within a section starting after skip rows.
            Stops at the next section header or after max_scan rows.

            Section-header heuristics (stops iteration):
              • Multi-word ALL-CAPS line with a space or comma  →  e.g. "VOIDS, RETURNS..."
              • Single ALL-CAPS word of 7+ alpha chars          →  e.g. "PAYMENTS", "REFUNDS"
            Short column headers like 'QTY', 'TOTAL', 'REASON', 'AMOUNT' (≤6 chars
            or mixed-case) do NOT stop iteration.
            """
            scanned = 0
            for r in range(start_row + skip, mr + 1):
                if scanned >= max_scan:
                    break
                lbl = s(r)
                if not lbl:
                    continue
                if lbl == lbl.upper() and len(lbl) > 5:
                    if ' ' in lbl or ',' in lbl:          # multi-word all-caps
                        break
                    if len(lbl) >= 7 and lbl.isalpha():   # single long all-caps word
                        break
                scanned += 1
                yield r

        result = dict(_ZERO)
        actual_cash = None
        expected_cash = None
        labor_rows_seen = set()

        # ── Pass 1: simple row-label lookups (no section context needed) ─────
        for row in range(1, mr + 1):
            lbl = s(row)

            # Sales handled separately below via SALES BY CLASS section

            if not result['transactions'] and 'Total Transactions' in lbl:
                v = num(row, 2)
                result['transactions'] = int(v) if v else 0

            if not result['avgTicket'] and 'Average Check' in lbl:
                result['avgTicket'] = num(row, 2) or 0

            # Labor is split across Management + Service Crew rows; sum both.
            # Only the first occurrence of each label counts, to avoid
            # double-counting if the label repeats elsewhere in the sheet.
            if lbl in LABOR_ROW_LABELS and lbl not in labor_rows_seen:
                labor_rows_seen.add(lbl)
                result['hoursWorked'] += num(row, 2) or 0
                result['laborCost'] += num(row, 3) or 0  # Wages column (usually column C)

            if 'Actual total cash to business' in lbl:
                actual_cash = num(row, 2) or 0

            if 'Expected total cash to business' in lbl:
                expected_cash = num(row, 2) or 0

            if not result['noSales'] and 'Number of No Sales' in lbl:
                v = num(row, 2)
                result['noSales'] = int(v) if v else 0

            if 'Pay Ins' in lbl and 'Pay Outs' not in lbl:
                result['payIns'] = num(row, 2) or 0

            if 'Pay Outs' in lbl:
                result['payOuts'] = num(row, 2) or 0

        if actual_cash is not None and expected_cash is not None:
            result['cashVariance'] = actual_cash - expected_cash

        # ── Sales: SALES BY CLASS section → "Total" row → col K (Net Totals) ──
        # The label in col A is just "Total" (not "Total Sales"), so a generic
        # label search would find it in many wrong places.  Anchor to the section.
        sales_section = find_section('SALES BY CLASS')
        if sales_section is not None:
            for r in range(sales_section + 1, sales_section + 10):
                if s(r).lower() == 'total':
                    v = num(r, 11)   # col K = Net Totals
                    if v:
                        result['sales'] = v
                    break

        # ── Pass 2: VOIDS section — sum Qty column ───────────────────────────
        vr = find_section('VOID')
        if vr is not None:
            qc = header_col(vr, {'qty', 'quantity', 'count', 'cnt', '#'})
            for r in section_rows(vr, skip=1):
                v = num(r, qc)
                if v:
                    result['clearedItems'] += v

        # ── Pass 3: Online sales — sum service sub-rows under "Custom Payment" ──
        # The PAYMENTS header row has column labels (Qty, Sales, Deposits, ...).
        # Detect which column is "Sales" from that header so position shifts don't
        # break extraction.  "Custom Payment" (row 64) is a formula-sum row whose
        # value isn't readable without data_only=True; the sub-rows below it
        # (DoorDash, Grubhub, etc.) are raw POS values — sum those instead.
        # Stop when "Grand Total" is reached to avoid double-counting.
        # Note: there are two "CUSTOM PAYMENT" sections (Tips and Payments);
        # we want the one under PAYMENTS, so search for it after pr_header.
        pr_header = find_section('PAYMENTS')
        online_sales_col = 3   # default: col C based on known report structure
        cr = None
        if pr_header is not None:
            for col in range(1, 20):
                if s(pr_header, col).lower() in ('sales', 'net sales', 'amount'):
                    online_sales_col = col
                    break
            # Find "CUSTOM PAYMENT" within the PAYMENTS section (after pr_header)
            kw = 'CUSTOM PAYMENT'.upper()
            for r in range(pr_header, mr + 1):
                if kw in s(r).upper():
                    cr = r
                    break
        if cr is not None:
            for r in section_rows(cr, skip=1):
                lbl = s(r)
                if 'total' in lbl.lower() or 'grand' in lbl.lower():
                    break   # hit the Grand Total row — stop
                v = num(r, online_sales_col)
                if v:
                    result['onlineSales'] += v

        # ── Pass 4: PAYMENTS section — sum Refunds column (col F) ────────────
        pr = find_section('PAYMENTS')
        if pr is not None:
            for r in section_rows(pr, skip=1):
                v = num(r, 6)
                if v:
                    result['refunds'] += abs(v)

        return result

    except Exception as e:
        print(f"Error extracting data: {e}")
        return dict(_ZERO)
